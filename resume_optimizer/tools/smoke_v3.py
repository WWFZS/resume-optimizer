# -*- coding: utf-8 -*-
"""v3 ASCII smoke self-check (no server required for static checks).

1) deleted endpoints absent from main.py source
2) editor_boxes merge logic (imports real main._merge_editor_boxes)
3) JS syntax: node --check on new/patched files
4) real export round-trip: resume JSON -> docx -> pdf -> long png
   + added-box injection into a real template docx

Run:  D:\\python\\python.exe tools\\smoke_v3.py
"""
import io
import re
import subprocess
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

OK, FAIL = "[PASS]", "[FAIL]"
results = []


def check(name, fn):
    try:
        fn()
        results.append((OK, name))
    except Exception as exc:  # noqa: BLE001
        results.append((FAIL, "%s -> %r" % (name, exc)))


# ── 1) deleted endpoints ──
def t_deleted_endpoints():
    src = (BASE / "main.py").read_text(encoding="utf-8")
    gone = [
        "/api/admin/user-vip",
        "/admin/templates/{template_id}/mapping",
        "/admin/aliases",
    ]
    for g in gone:
        assert g not in src, "still present: " + g
    for need in ["/api/export/{fmt}", "_merge_editor_boxes", "/api/template-save/{template_id}",
                 "/api/template-preview/{template_id}", "/api/draft/save", "/api/admin/template-config/{template_id}"]:
        assert need in src, "missing: " + need


# ── 2) structure merge ──
def t_merge():
    from main import _merge_editor_boxes
    st = {"page_w": 794, "page_h": 1123, "boxes": [
        {"id": 0, "kind": "textbox", "x": 10, "y": 10, "w": 100, "h": 30, "text": "a"},
        {"id": 1, "kind": "table", "x": 20, "y": 40, "w": 200, "h": 60, "text": "b"},
        {"id": 2, "kind": "textbox", "x": 30, "y": 90, "w": 100, "h": 30, "text": "c"},
    ]}
    _merge_editor_boxes(st, {
        "overrides": {"1": {"x": 50, "y": 60, "w": 180, "h": 55, "editable": False},
                      "2": {"deleted": True}},
        "added": [{"id": -1, "x": 60, "y": 60, "w": 200, "h": 40, "text": "new",
                   "font": {"family": "微软雅黑", "size": 14, "color": "#333333"}, "align": "center"}],
    })
    ids = [b["id"] for b in st["boxes"]]
    assert ids == [0, 1, -1], ids
    b1 = st["boxes"][1]
    assert (b1["x"], b1["y"], b1["w"], b1["h"], b1["editable"]) == (50, 60, 180, 55, False)
    added = st["boxes"][2]
    assert added["added"] and added["text"] == "new" and added["align"] == "center"


# ── 3) JS syntax ──
def t_js():
    files = ["static/editor-v3.js", "static/resume-render.js",
             "static/resume-bind.js", "static/editor-richtext.js"]
    for f in files:
        r = subprocess.run(["node", "--check", str(BASE / f)], capture_output=True, text=True)
        assert r.returncode == 0, "%s: %s" % (f, r.stderr[:200])
    for html in ["static/editor.html", "static/template-editor.html", "static/admin.html"]:
        src = (BASE / html).read_text(encoding="utf-8")
        m = re.search(r"<script>(.*)</script>", src, re.S)
        if m:
            tmp = BASE / "temp" / "_smoke_inline.js"
            tmp.write_text(m.group(1), encoding="utf-8")
            r = subprocess.run(["node", "--check", str(tmp)], capture_output=True, text=True)
            assert r.returncode == 0, "%s inline: %s" % (html, r.stderr[:200])
    ed = (BASE / "static/editor.html").read_text(encoding="utf-8")
    assert "editor.js" not in ed and "editor.css" not in ed
    assert "editor-v3.js" in ed and "resume-render.js" in ed
    # 真实模板「文档直接编辑」模式已接线：模块 + 样式 + 桥接挂载点齐备
    assert "editor-richtext.js" in ed, "editor.html 未加载 editor-richtext.js"
    assert "editor-richtext.css" in ed, "editor.html 未加载 editor-richtext.css"
    assert 'id="previewPaper"' in ed, "editor.html 缺少富文本纸面 #previewPaper"
    assert 'id="moduleList"' in ed, "editor.html 缺少文档导航 #moduleList"
    v3 = (BASE / "static/editor-v3.js").read_text(encoding="utf-8")
    assert "__edBridge" in v3, "editor-v3.js 未暴露 __edBridge（richtext 依赖）"
    assert "richtext_mode" in v3, "editor-v3.js 未在草稿中持久化 richtext_mode"
    rtjs = (BASE / "static/editor-richtext.js").read_text(encoding="utf-8")
    assert "__edBridge" in rtjs and "RichtextMode" in rtjs
    # B 阶段：槽位映射模块已接线（seed + writeback 两个钩子都在）
    assert "resume-bind.js" in ed, "editor.html 未加载 resume-bind.js"
    rb = (BASE / "static/resume-bind.js")
    assert rb.exists(), "static/resume-bind.js 缺失"
    assert "seedFromResume()" in rtjs, "richtext 未接 seed 钩子"
    assert "applyBoxEdit" in rtjs, "richtext 未接 writeback 钩子"
    admin = (BASE / "static/admin.html").read_text(encoding="utf-8")
    for token in ["vipUserCard", "mappingCard", "aliasCard", "cfgModal", "user-vip"]:
        assert token not in admin, "admin.html still has " + token
    assert "template-editor.html?id=" in admin


# ── 4) export round-trip ──
def t_export_roundtrip():
    from utils.resume_docx import generate_docx, pdf_to_long_png
    from utils.converter import docx_to_pdf
    resume = {
        "basic": {"name": "Smoke Tester", "phone": "13800000000", "email": "s@t.com",
                  "city": "Beijing", "intention": "Backend Engineer",
                  "photo": "" if True else ""},
        "modules": {
            "education": [{"school": "Test Univ", "degree": "B.Sc.", "major": "CS",
                           "start": "2016-09", "end": "2020-06",
                           "content": "line one\nline two"}],
            "work": [{"company": "Acme", "position": "Engineer", "start": "2020-07",
                      "current": True, "content": "built stuff\nshipped things"}],
            "skill": [{"name": "Python", "level": "expert", "content": "FastAPI"}],
            "honor": [{"time": "2021", "name": "Best Award", "issuer": "Org"}],
            "self": "self evaluation text",
            "hobby": ["reading", "chess"],
        },
        "order": ["education", "work", "skill", "honor", "self", "hobby"],
        "hidden": [], "showPhoto": False,
    }
    theme = {"family": "single", "primary": "#1E5AE8", "font": "yahei", "fontSize": 13,
             "lineHeight": 1.72, "gap": 20, "nameSize": 30, "photoShape": "rect",
             "photoSize": 92, "padding": 44, "sectionStyle": "underline"}
    docx_bytes = generate_docx(resume, theme)
    assert len(docx_bytes) > 10000
    docx_path = BASE / "temp" / "_smoke_v3.docx"
    pdf_path = BASE / "temp" / "_smoke_v3.pdf"
    docx_path.write_bytes(docx_bytes)
    docx_to_pdf(str(docx_path), str(pdf_path))
    assert pdf_path.read_bytes()[:4] == b"%PDF"
    png = pdf_to_long_png(str(pdf_path))
    assert png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) > 20000


# ── 5) added-box injection on real template ──
def t_added_injection():
    from utils import template_editor
    tpl = BASE / "temp" / "preview_regenerate" / "t1.docx"
    if not tpl.exists():
        return  # skip silently when sample template missing
    added = [{"id": -1, "x": 60, "y": 600, "w": 300, "h": 60,
              "text": "SMOKE-INJECT-OK", "align": "center",
              "font": {"family": "微软雅黑", "size": 16, "color": "#1E5AE8", "bold": True}}]
    out = template_editor.apply_box_texts(str(tpl), {}, added)
    assert len(out) > 10000
    buf = io.BytesIO(out)
    from docx import Document
    xml = Document(buf).element.body.xml
    assert "SMOKE-INJECT-OK" in xml and "AddedBox-1" in xml


# ── 6) resume-render accent mix ──
def t_mix_theme():
    import subprocess
    js = ("global.window={};const mod=require('%s');"
          "const t=window.ResumeRender.mixTheme('#1E5AE8');"
          "console.log(JSON.stringify([t.accent,t.deep.length,t.soft.length===7]));"
          ) % str(BASE / "static" / "resume-render.js").replace("\\", "/")
    r = subprocess.run(["node", "-e", js], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr[:200]
    assert '"#1E5AE8"' in r.stdout and "true" in r.stdout, r.stdout


# ── 7) resume-bind seed/writeback round-trip (offline fixture) ──
def t_resume_bind_roundtrip():
    js = r"""
const RB = require('%(RB)s');
const structure = { boxes: [
  { id: 0, kind: "free", role: null, y: 40, placeholders: ["求职意向"], text: "简历模板资源网\n求职意向：{{求职意向}}" },
  { id: 1, kind: "field", role: "basic", y: 60, placeholders: ["姓名"], text: "{{姓名}}" },
  { id: 2, kind: "content", role: "education", y: 200, placeholders: ["工作经历"], text: "2014年寒假 肯德基 服务员" },
  { id: 3, kind: "title", role: "education", y: 190, placeholders: [], text: "教育背景" },
  { id: 4, kind: "content", role: "basic", y: 240, placeholders: ["教育背景"], text: "{{教育背景}}\n主修课程：样例课" },
  { id: 6, kind: "field", role: "basic", y: 120, placeholders: ["所在地", "电话", "邮箱"],
    text: "生日：1994.04.06\n性别：男\n现居：{{所在地}}\n电话：{{电话}}\n邮箱：{{邮箱}}" },
  { id: 8, kind: "content", role: "internship", y: 320, placeholders: [], text: "2016年暑假 某银行 大堂经理" }
]};
const resume = { basic: { name: "测试员", intention: "数据分析师", phone: "13800000000",
    email: "t@e.com", city: "贵阳", gender: "男", birth_date: "1995-01-02" },
  modules: {
    education: [{ school: "测试大学", major: "计算机", degree: "本科", start: "2020-09", end: "2024-06", content: "主修：数据结构" }],
    internship: [{ company: "大华科技", position: "数据实习生", start: "2025-01", end: "2025-06", content: "看板搭建" }],
    work: [{ company: "某某银行", position: "大堂经理", start: "2016-07", end: "2017-09", content: "客户服务" }],
    project: [], campus: [], skill: [], honor: [], self: "细致耐心。", custom: "", hobby: [] } };
const fails = [];
const assert = (cond, msg) => { if (!cond) fails.push(msg); };
const txt = e => e && e.paragraphs.map(p => p.runs.map(r => r.text).join("")).join("\n");

/* seed: role 语义优先于误标 role（box2 占位符=工作经历 → work 记录） */
const edits = RB.seedEdits(structure, resume);
assert(txt(edits[2]).includes("某某银行"), "box2 work seed");
assert(txt(edits[2]).includes("2016-07 - 2017-09"), "box2 period");
assert(txt(edits[4]).includes("测试大学"), "box4 education seed");
assert(!txt(edits[4]).includes("样例课"), "box4 sample purged");
assert(txt(edits[1]) === "测试员", "box1 name");
assert(txt(edits[6]).includes("1995-01-02") && txt(edits[6]).includes("13800000000"), "box6 labels");
assert(txt(edits[6]).indexOf("1994") < 0, "box6 sample birthday purged");
assert(txt(edits[8]).includes("大华科技"), "box8 internship by declared role");
assert(txt(edits[0]).includes("数据分析师"), "box0 intention");

/* writeback: 编辑盒子 → 结构化字段，且解析不丢整段 */
const r2 = JSON.parse(JSON.stringify(resume));
assert(RB.applyBoxEdit(structure, r2, 8, "2025年3月 - 2025年8月    新公司    新岗位\n改过的描述"), "wb8 ok");
assert(r2.modules.internship[0].company === "新公司", "wb8 company: " + JSON.stringify(r2.modules.internship[0]));
assert(r2.modules.internship[0].content === "改过的描述", "wb8 content");
assert(r2.modules.internship[0].start === "2025-03", "wb8 start");
assert(RB.applyBoxEdit(structure, r2, 6, "生日：1996.02.03\n性别：女\n现居：上海\n电话：13911112222\n邮箱：a@b.c"), "wb6 ok");
assert(r2.basic.city === "上海" && r2.basic.phone === "13911112222" && r2.basic.gender === "女", "wb6 basic");
/* 换模板 → 新结构 seed 出同样内容（内容跟模板走，与 box.id 无关） */
const edits2 = RB.seedEdits(structure, r2);
assert(txt(edits2[8]).includes("新公司"), "reseed follows writeback");
/* coverage */
const cov = RB.coverage(structure, resume);
assert(cov.work && cov.work.records === 1 && cov.work.boxes === 1, "coverage work");
console.log(fails.length ? "FAILS: " + fails.join(" | ") : "ROUNDTRIP-OK");
process.exit(fails.length ? 1 : 0);
""" % {"RB": str(BASE / "static" / "resume-bind.js").replace("\\", "/")}
    r = subprocess.run(["node", "-e", js], capture_output=True, text=True)
    assert r.returncode == 0, (r.stdout + r.stderr)[:400]
    assert "ROUNDTRIP-OK" in r.stdout, r.stdout[:400]


for name, fn in [
    ("deleted endpoints + kept routes in main.py", t_deleted_endpoints),
    ("_merge_editor_boxes overrides/added", t_merge),
    ("JS syntax (node --check) + reference hygiene", t_js),
    ("export round-trip docx->pdf->long png", t_export_roundtrip),
    ("added-box injection into real template", t_added_injection),
    ("resume-render mixTheme accent", t_mix_theme),
    ("resume-bind seed/writeback round-trip", t_resume_bind_roundtrip),
]:
    check(name, fn)

print("=" * 62)
for mark, name in results:
    print("%s %s" % (mark, name))
fails = [r for r in results if r[0] == FAIL]
print("=" * 62)
print("SMOKE RESULT: %d passed, %d failed" % (len(results) - len(fails), len(fails)))
sys.exit(1 if fails else 0)
